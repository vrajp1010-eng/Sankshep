import io
import os
import re
import asyncio
import logging
import docx
import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
from typing import Optional, List, Dict, Any
from dotenv import load_dotenv

from chains import get_chain, build_input_vars, resolve_transformation_type, TEMPLATES
from db import init_db, log_transformation, get_recent_history, get_analytics_summary
from providers import (
    PROVIDER_REGISTRY,
    list_public_providers,
    validate_provider_id,
    get_provider_model,
    is_provider_available,
    check_provider_startup_status,
    redact_secrets,
)

load_dotenv()

# Configure logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("sankshep.api")

app = FastAPI(
    title="Sankshep.ai - AI Content Transformation Platform",
    description="Multi-provider AI content transformation and comparison engine.",
    version="2.0.0",
)

init_db()
check_provider_startup_status()

MAX_FILE_SIZE_MB = 15

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class TransformRequest(BaseModel):
    text: str
    transformation_type: str
    provider: Optional[str] = None
    target_language: Optional[str] = None
    tone: Optional[str] = None
    brand_voice_name: Optional[str] = None
    brand_voice_notes: Optional[str] = None


class CompareRequest(BaseModel):
    text: str
    transformation_type: str
    providers: List[str] = Field(..., description="List of provider IDs to compare")
    target_language: Optional[str] = None
    tone: Optional[str] = None
    brand_voice_name: Optional[str] = None
    brand_voice_notes: Optional[str] = None


class UrlExtractRequest(BaseModel):
    url: str


class PptxExportRequest(BaseModel):
    slide_text: str


# ── /providers ───────────────────────────────────────────────────────────────

@app.get("/providers")
async def get_providers():
    """
    Returns the list of available and supported AI providers without exposing
    any secrets, keys, or internal environment configurations.
    """
    return {"providers": list_public_providers()}


# ── /status ──────────────────────────────────────────────────────────────────

@app.get("/status")
async def status_api():
    """
    Health check and active provider telemetry.
    """
    public_providers = list_public_providers()
    available_providers = [p for p in public_providers if p["available"]]

    active_provider = available_providers[0] if available_providers else None

    return {
        "status": "online",
        "name": "Sankshep.ai",
        "engine": active_provider["label"] if active_provider else "None",
        "model": active_provider["model"] if active_provider else "None",
        "available_providers": [p["id"] for p in available_providers],
        "ready": len(available_providers) > 0,
    }


# ── /transform ────────────────────────────────────────────────────────────────

@app.post("/transform")
async def transform_api(req: TransformRequest):
    """
    Single-provider token-by-token streaming transformation.
    """
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Source text cannot be empty.")

    # Validate provider (400 if unknown, 503 if unconfigured)
    resolved_provider = validate_provider_id(req.provider)

    try:
        chain = get_chain(req.transformation_type, provider_id=resolved_provider, streaming=True)
        input_vars = build_input_vars(
            text=req.text,
            transformation_type=req.transformation_type,
            target_language=req.target_language,
            tone=req.tone,
            brand_voice_name=req.brand_voice_name,
            brand_voice_notes=req.brand_voice_notes,
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except HTTPException:
        raise
    except Exception as e:
        safe_msg = redact_secrets(str(e))
        logger.error(f"Error initializing chain: {safe_msg}")
        raise HTTPException(status_code=500, detail="Failed to initialize AI transformation chain.")

    input_word_count = len(req.text.split())

    def stream_and_log():
        collected = ""
        try:
            for chunk in chain.stream(input_vars):
                collected += chunk
                yield chunk
        except Exception as e:
            safe_err = redact_secrets(str(e))
            logger.error(f"Streaming error for provider '{resolved_provider}': {safe_err}")
            yield f"\n\n[Generation Error: Provider request failed. Please try another provider.]"
        finally:
            output_word_count = len(collected.split())
            if output_word_count > 0:
                log_transformation(req.transformation_type, input_word_count, output_word_count)

    return StreamingResponse(stream_and_log(), media_type="text/plain; charset=utf-8")


# ── /transform/compare ────────────────────────────────────────────────────────

async def _run_single_provider_generation(
    provider_id: str,
    text: str,
    transformation_type: str,
    target_language: Optional[str],
    tone: Optional[str],
    brand_voice_name: Optional[str],
    brand_voice_notes: Optional[str],
    timeout_seconds: float = 35.0,
) -> Dict[str, Any]:
    """
    Runs generation for a single provider inside comparison mode.
    Guarantees isolated execution with individual timeout and safe error reporting.
    """
    entry = PROVIDER_REGISTRY.get(provider_id)
    label = entry["label"] if entry else provider_id
    model = get_provider_model(provider_id)

    if not is_provider_available(provider_id):
        return {
            "provider": provider_id,
            "label": label,
            "model": model,
            "status": "error",
            "error": f"Provider {label} is not configured on the server. Please try another provider.",
        }

    try:
        # Run blocking LLM invocation in threadpool with timeout
        def _invoke():
            chain = get_chain(transformation_type, provider_id=provider_id, streaming=False)
            input_vars = build_input_vars(
                text=text,
                transformation_type=transformation_type,
                target_language=target_language,
                tone=tone,
                brand_voice_name=brand_voice_name,
                brand_voice_notes=brand_voice_notes,
            )
            return chain.invoke(input_vars)

        output = await asyncio.wait_for(asyncio.to_thread(_invoke), timeout=timeout_seconds)

        # Log metrics to DB
        input_word_count = len(text.split())
        output_word_count = len(str(output).split())
        log_transformation(f"{transformation_type}:{provider_id}", input_word_count, output_word_count)

        return {
            "provider": provider_id,
            "label": label,
            "model": model,
            "status": "ok",
            "output": str(output),
        }
    except asyncio.TimeoutError:
        logger.warning(f"Provider {provider_id} timed out after {timeout_seconds}s")
        return {
            "provider": provider_id,
            "label": label,
            "model": model,
            "status": "error",
            "error": f"Provider {label} request timed out. Please try another provider.",
        }
    except Exception as e:
        safe_msg = redact_secrets(str(e))
        logger.error(f"Provider {provider_id} failed: {safe_msg}")
        return {
            "provider": provider_id,
            "label": label,
            "model": model,
            "status": "error",
            "error": f"Provider {label} request failed. Please try another provider.",
        }


@app.post("/transform/compare")
async def compare_api(req: CompareRequest):
    """
    Executes multiple providers concurrently for side-by-side comparison.
    Deduplicates IDs, validates min 2, runs independently with timeouts.
    """
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Source text cannot be empty.")

    resolved_type = resolve_transformation_type(req.transformation_type)
    if resolved_type not in TEMPLATES:
        raise HTTPException(status_code=400, detail=f"Unsupported transformation type: '{req.transformation_type}'.")

    # Deduplicate provider IDs while preserving order
    deduped_ids = []
    seen = set()
    for pid in req.providers:
        cleaned = str(pid).strip().lower()
        if cleaned not in seen:
            seen.add(cleaned)
            deduped_ids.append(cleaned)

    # Validate provider count: require at least 2
    if len(deduped_ids) < 2:
        raise HTTPException(status_code=400, detail="Comparison mode requires at least 2 distinct providers.")

    # Cap at registry size
    max_providers = len(PROVIDER_REGISTRY)
    if len(deduped_ids) > max_providers:
        raise HTTPException(status_code=400, detail=f"Cannot compare more than {max_providers} providers.")

    # Validate that all requested provider IDs exist in allowlist
    for pid in deduped_ids:
        if pid not in PROVIDER_REGISTRY:
            allowed = list(PROVIDER_REGISTRY.keys())
            raise HTTPException(status_code=400, detail=f"Unknown provider '{pid}'. Allowed providers: {allowed}")

    # Launch all providers concurrently with independent error isolation
    tasks = [
        _run_single_provider_generation(
            provider_id=pid,
            text=req.text,
            transformation_type=req.transformation_type,
            target_language=req.target_language,
            tone=req.tone,
            brand_voice_name=req.brand_voice_name,
            brand_voice_notes=req.brand_voice_notes,
        )
        for pid in deduped_ids
    ]

    results = await asyncio.gather(*tasks)

    return {"results": results}


# ── /extract/url ──────────────────────────────────────────────────────────────

@app.post("/extract/url")
async def extract_url(req: UrlExtractRequest):
    url = req.url.strip()
    if not url:
        raise HTTPException(status_code=400, detail="URL cannot be empty.")
    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }

    try:
        resp = requests.get(url, headers=headers, timeout=12)
        resp.raise_for_status()
        
        soup = BeautifulSoup(resp.content, "html.parser")
        
        # Remove noisy elements
        for element in soup(["script", "style", "nav", "footer", "header", "noscript", "aside", "svg", "form"]):
            element.extract()
            
        # Extract title
        title = soup.title.string.strip() if soup.title and soup.title.string else ""
        
        # Find main content container if possible
        main_content = soup.find("article") or soup.find("main") or soup.find("div", class_=re.compile(r"content|post|article|body", re.I)) or soup.body
        
        if not main_content:
            main_content = soup
            
        text = main_content.get_text(separator="\n", strip=True)
        # Clean multiple blank lines
        clean_text = re.sub(r"\n{3,}", "\n\n", text)
        
        if len(clean_text) < 50:
            raise HTTPException(status_code=422, detail="Could not extract sufficient text content from this webpage.")
            
        return {
            "title": title,
            "url": url,
            "extracted_text": clean_text[:25000],  # Limit to 25k chars
        }
    except requests.exceptions.RequestException as e:
        raise HTTPException(status_code=400, detail=f"Failed to fetch webpage: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error parsing webpage: {str(e)}")


# ── /upload ───────────────────────────────────────────────────────────────────

@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = file.filename.lower()
    allowed_exts = (".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".png", ".jpg", ".jpeg", ".webp")
    
    if not any(filename.endswith(ext) for ext in allowed_exts):
        raise HTTPException(
            status_code=400,
            detail="Supported formats: PDF, DOCX, TXT, MD, CSV, JSON, PNG, JPG, WEBP.",
        )

    content = await file.read()

    if len(content) > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail=f"File exceeds maximum allowed size of {MAX_FILE_SIZE_MB}MB.",
        )

    text = ""
    try:
        if filename.endswith(".pdf"):
            pdf = PdfReader(io.BytesIO(content))
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"

            # Fallback to OCR if PDF has no text layer
            if not text.strip():
                text = _ocr_pdf(content)

        elif filename.endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

        elif filename.endswith((".txt", ".md", ".csv", ".json")):
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1", errors="ignore")

        elif filename.endswith((".png", ".jpg", ".jpeg", ".webp")):
            text = _ocr_image(content)

        if not text.strip():
            raise HTTPException(
                status_code=400,
                detail="No readable text could be extracted from the uploaded document.",
            )

        return {"filename": file.filename, "extracted_text": text.strip()}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process file: {str(e)}")


def _ocr_pdf(content: bytes) -> str:
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except ImportError:
        raise HTTPException(
            status_code=422,
            detail="Scanned image PDF detected. Please install pytesseract and pdf2image for OCR support.",
        )

    try:
        images = convert_from_bytes(content, dpi=200)
        ocr_text = ""
        for img in images:
            ocr_text += pytesseract.image_to_string(img) + "\n"
        return ocr_text
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"OCR processing failed: {str(e)}.",
        )


def _ocr_image(content: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        raise HTTPException(
            status_code=422,
            detail="Image OCR requires pytesseract and Pillow.",
        )

    try:
        image = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(image)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Image OCR failed: {str(e)}")


# ── /export/pptx ─────────────────────────────────────────────────────────────

def _parse_slides(slide_text: str) -> list:
    slides = []
    current_title = None
    current_bullets = []

    for line in slide_text.strip().splitlines():
        line = line.strip()
        if not line:
            continue

        slide_match = re.match(r"^(?:#+\s*|\*\*\s*)?Slide\s+\d+[:.]\s*(.*?)(?:\*\*)?$", line, re.IGNORECASE)
        if slide_match:
            if current_title is not None:
                slides.append({"title": current_title, "bullets": current_bullets})
            current_title = slide_match.group(1).strip()
            current_bullets = []
        elif line.startswith(("-", "•", "*")):
            current_bullets.append(line.lstrip("-•* ").strip())
        elif current_title is not None:
            current_bullets.append(line)

    if current_title is not None:
        slides.append({"title": current_title, "bullets": current_bullets})

    return slides


@app.post("/export/pptx")
async def export_pptx(req: PptxExportRequest):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise HTTPException(
            status_code=501,
            detail="python-pptx is not installed.",
        )

    slides_data = _parse_slides(req.slide_text)
    if not slides_data:
        raise HTTPException(
            status_code=400,
            detail="Could not parse slide structures from text. Ensure slides use 'Slide 1: Title' format.",
        )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ── Title Slide ──────────────────────────────────────────────────────────
    title_layout = prs.slide_layouts[0]
    cover = prs.slides.add_slide(title_layout)
    cover.shapes.title.text = "AI-Generated Presentation Deck"
    subtitle_ph = cover.placeholders[1]
    if subtitle_ph:
        subtitle_ph.text = "Generated by Sankshep.ai"

    # ── Content Slides ───────────────────────────────────────────────────────
    content_layout = prs.slide_layouts[1]
    for slide_data in slides_data:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = slide_data["title"]

        tf = slide.placeholders[1].text_frame
        tf.clear()

        for idx, bullet in enumerate(slide_data["bullets"]):
            if idx == 0:
                tf.paragraphs[0].text = bullet
                tf.paragraphs[0].level = 0
            else:
                p = tf.add_paragraph()
                p.text = bullet
                p.level = 0

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": 'attachment; filename="Sankshep_Presentation.pptx"'},
    )


# ── /history & /analytics ───────────────────────────────────────────────────

@app.get("/history")
async def history_api(limit: int = 25):
    return {"history": get_recent_history(limit)}


@app.get("/analytics")
async def analytics_api():
    return get_analytics_summary()